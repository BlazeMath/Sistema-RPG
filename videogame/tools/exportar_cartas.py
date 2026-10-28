"""Exportador de cartas do Sistema 21.

Lê a planilha oficial de cartas (abas Classes, Raças e Backgrounds) e gera
os arquivos JSON que o jogo usa:

    data/cartas/habilidades.json   cartas de habilidade de dois lados
    data/cartas/herois.json        metades de carta de herói
    data/cartas/relatorio.txt      avisos de validação e mudanças desde a última exportação

Uso:
    python tools/exportar_cartas.py caminho/para/SISTEMA_21_Cartas_Oficial.xlsx

Cada carta recebe uma "impressão digital" (hash do texto). Quando uma carta muda
na planilha, o relatório aponta qual carta mudou para a implementação ser revista.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ENTIDADES = {
    "Classes": ["GUERREIRO", "LADINO", "CLÉRIGO", "CAÇADOR", "MAGO"],
    "Raças": ["HUMANO", "ELFO", "ANÃO", "HALFLING", "ORC"],
    "Backgrounds": ["SOLDADO", "BANDIDO", "ESCRAVO", "EXPLORADOR", "FERVOROSO"],
}
TIPO_ENTIDADE = {"Classes": "classe", "Raças": "raca", "Backgrounds": "background"}

CAMPOS_HABILIDADE = {"Nome": "nome", "Tags": "tags", "Descrição": "texto", "Limite fad.": "limite_fadiga"}
CAMPOS_HEROI = {
    "Passiva": "passiva",
    "Ativada": "ativada",
    "Desencadeada": "desencadeada",
    "Ao Critar 21": "critico",
    "Crítico": "critico",
    "Limite fad.": "limite_fadiga",
}

# Tipos de ação reconhecidos no custo (antes dos dois-pontos), do mais longo ao mais curto.
TIPOS_ACAO = [
    ("ação de movimento", "movimento"),
    ("ação padrão", "padrao"),
    ("ação livre", "livre"),
    ("reação", "reacao"),
    ("ação", "padrao"),
]


def slug(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^a-zA-Z0-9]+", "-", texto).strip("-").lower()
    return texto


def limpar(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def ler_custo(texto: str | None) -> dict:
    """Separa custo e efeito na convenção 'custo: efeito'.

    A linha de custo é a primeira linha cujo trecho antes dos dois-pontos fala de
    fadiga, ação, reação, descarte, exílio ou fichas. Linhas anteriores a ela são
    efeitos permanentes da carta (ex.: "Concede Fogo."). Sem linha de custo, a
    habilidade é passiva ou desencadeada e o texto inteiro é efeito.
    """
    resultado = {"fadiga": 0, "acao": None, "outros": [], "permanente": "", "efeito": texto or ""}
    if not texto:
        return resultado
    linhas = texto.split("\n")
    for i, linha in enumerate(linhas):
        if ":" not in linha:
            continue
        antes, depois = linha.split(":", 1)
        if not re.search(r"fadiga|ação|reação|descarte|exile|ficha|mira", antes.lower()):
            continue
        for parte in [p.strip() for p in antes.split(",") if p.strip()]:
            parte_min = parte.lower()
            m = re.fullmatch(r"(\d+)\s+de\s+fadiga(?:\s*\((.+)\))?", parte_min)
            if m:
                resultado["fadiga"] = int(m.group(1))
                parte_min = (m.group(2) or "").strip()
                if not parte_min:
                    continue
            tipo = next((t for chave, t in TIPOS_ACAO if parte_min == chave), None)
            if tipo:
                resultado["acao"] = tipo
                continue
            resultado["outros"].append(parte)
        resultado["permanente"] = "\n".join(linhas[:i]).strip()
        resultado["efeito"] = "\n".join([depois.strip(), *linhas[i + 1:]]).strip()
        return resultado
    return resultado


def impressao_digital(dados: dict) -> str:
    bruto = json.dumps(dados, ensure_ascii=False, sort_keys=True)
    return hashlib.sha1(bruto.encode("utf-8")).hexdigest()[:12]


def para_int(valor, avisos: list, contexto: str) -> int | None:
    if valor is None:
        avisos.append(f"{contexto}: limite de fadiga vazio")
        return None
    try:
        return int(str(valor).strip())
    except ValueError:
        avisos.append(f"{contexto}: limite de fadiga não numérico ({valor!r})")
        return None


def exportar_aba(ws, nome_aba: str, avisos: list):
    entidades = ENTIDADES[nome_aba]
    tipo = TIPO_ENTIDADE[nome_aba]
    habilidades, herois = [], []
    entidade = grupo = None
    atual = None  # carta ou herói em montagem
    esperando_sabor = False

    def fechar():
        nonlocal atual
        if atual is None:
            return
        (herois if atual["_tipo"] == "heroi" else habilidades).append(atual)
        atual = None

    for linha in ws.iter_rows(values_only=True):
        a = limpar(linha[0]) if len(linha) > 0 else None
        b = limpar(linha[1]) if len(linha) > 1 else None
        c = limpar(linha[2]) if len(linha) > 2 else None

        if a is None and b is None and c is None:
            continue
        if a in entidades and b is None:
            fechar()
            entidade, grupo = a, None
            continue
        if a is None and b == "LADO A":
            fechar()
            atual = {"_tipo": "habilidade", "grupo": grupo, "entidade": entidade, "lados": [{}, {}]}
            continue
        if a and a.startswith("★"):
            fechar()
            atual = {"_tipo": "heroi", "grupo": grupo, "entidade": entidade, "nome": a.lstrip("★ ").strip()}
            esperando_sabor = True
            continue
        if atual and atual["_tipo"] == "heroi" and esperando_sabor and b is None and a not in CAMPOS_HEROI:
            atual["descricao"] = a
            esperando_sabor = False
            continue
        if atual and atual["_tipo"] == "heroi" and a in CAMPOS_HEROI:
            esperando_sabor = False
            atual[CAMPOS_HEROI[a]] = b
            continue
        if atual and atual["_tipo"] == "habilidade" and a in CAMPOS_HABILIDADE:
            for i, valor in enumerate((b, c)):
                atual["lados"][i][CAMPOS_HABILIDADE[a]] = valor
            continue
        if b is None and a:
            # Cabeçalho de grupo (estratégia, meia cartas, 2 XP, heróis, aprendiz...).
            fechar()
            grupo = a
            continue
        avisos.append(f"{nome_aba}: linha não reconhecida: {a!r} | {b!r}")
    fechar()

    saida_hab, saida_her = [], []
    for carta in habilidades:
        entidade = carta["entidade"]
        lados = []
        for i, lado in enumerate(carta["lados"]):
            if not any(lado.values()):
                lados.append(None)
                continue
            nome = lado.get("nome") or "(sem nome)"
            ctx = f"{nome_aba}/{entidade}/{nome} lado {'AB'[i]}"
            if not lado.get("nome"):
                avisos.append(f"{ctx}: sem nome")
            if not lado.get("texto"):
                avisos.append(f"{ctx}: sem descrição")
            tags = [t.strip() for t in (lado.get("tags") or "").split(",") if t.strip()]
            lados.append({
                "nome": lado.get("nome"),
                "tags": tags,
                "texto": lado.get("texto"),
                "limite_fadiga": para_int(lado.get("limite_fadiga"), avisos, ctx),
                "custo": ler_custo(lado.get("texto")),
            })
        nome_a = (lados[0] or {}).get("nome") or "sem-nome"
        grupo_txt = carta["grupo"] or ""
        registro = {
            "id": f"{tipo}.{slug(entidade)}.{slug(nome_a)}",
            "fonte": tipo,
            "entidade": entidade.title(),
            "grupo": grupo_txt,
            "xp": 2 if "2 XP" in grupo_txt else 1,
            "meia_carta": grupo_txt.upper().startswith("MEIA"),
            "lados": lados,
        }
        registro["impressao"] = impressao_digital({"lados": lados})
        saida_hab.append(registro)

    for heroi in herois:
        entidade = heroi["entidade"]
        ctx = f"{nome_aba}/{entidade}/herói {heroi['nome']}"
        campos = {k: heroi.get(k) for k in ("passiva", "ativada", "desencadeada", "critico")}
        for k, v in campos.items():
            if not v:
                avisos.append(f"{ctx}: efeito '{k}' vazio")
        registro = {
            "id": f"{tipo}.{slug(entidade)}.heroi.{slug(heroi['nome'])}",
            "fonte": tipo,
            "entidade": entidade.title(),
            "nome": heroi["nome"].title() if heroi["nome"].isupper() else heroi["nome"],
            "descricao": heroi.get("descricao"),
            **campos,
            "custo_ativada": ler_custo(campos["ativada"]),
            "limite_fadiga": para_int(heroi.get("limite_fadiga"), avisos, ctx),
        }
        registro["impressao"] = impressao_digital({k: registro[k] for k in ("passiva", "ativada", "desencadeada", "critico", "limite_fadiga")})
        saida_her.append(registro)
    return saida_hab, saida_her


def corrigir_entidades_mago(habilidades: list) -> None:
    """No Mago, cada 'APRENDIZ DE <ELEMENTO>' é um grupo; marca o elemento."""
    for carta in habilidades:
        grupo = carta["grupo"].upper()
        if grupo.startswith("APRENDIZ DE "):
            carta["elemento"] = grupo.replace("APRENDIZ DE ", "").title()


def comparar(anteriores: list, novos: list, rotulo: str) -> list:
    antigos = {c["id"]: c["impressao"] for c in anteriores}
    linhas = []
    for c in novos:
        if c["id"] not in antigos:
            linhas.append(f"NOVA     {rotulo}: {c['id']}")
        elif antigos[c["id"]] != c["impressao"]:
            linhas.append(f"MUDOU    {rotulo}: {c['id']}")
    ids_novos = {c["id"] for c in novos}
    for cid in antigos:
        if cid not in ids_novos:
            linhas.append(f"REMOVIDA {rotulo}: {cid}")
    return linhas


def main(caminho_xlsx: str, pasta_saida: str = "data/cartas") -> int:
    saida = Path(pasta_saida)
    saida.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(caminho_xlsx, read_only=True)
    avisos: list[str] = []
    habilidades, herois = [], []
    for aba in ENTIDADES:
        if aba not in wb.sheetnames:
            avisos.append(f"aba '{aba}' não encontrada")
            continue
        h, he = exportar_aba(wb[aba], aba, avisos)
        habilidades += h
        herois += he
    corrigir_entidades_mago(habilidades)

    for lista, rotulo in ((habilidades, "habilidade"), (herois, "herói")):
        vistos = {}
        for c in lista:
            if c["id"] in vistos:
                vistos[c["id"]] += 1
                c["id"] = f"{c['id']}-{vistos[c['id']]}"
                avisos.append(f"id de {rotulo} repetido, renomeado para {c['id']}")
            else:
                vistos[c["id"]] = 1

    mudancas = []
    for nome, novos, rotulo in (("habilidades.json", habilidades, "habilidade"), ("herois.json", herois, "herói")):
        arquivo = saida / nome
        if arquivo.exists():
            mudancas += comparar(json.loads(arquivo.read_text(encoding="utf-8")), novos, rotulo)
        arquivo.write_text(json.dumps(novos, ensure_ascii=False, indent=1), encoding="utf-8")

    relatorio = [
        f"Planilha: {Path(caminho_xlsx).name}",
        f"Cartas de habilidade: {len(habilidades)}",
        f"Metades de herói: {len(herois)}",
        "",
        f"Avisos ({len(avisos)}):",
        *[f"  - {a}" for a in avisos],
        "",
        f"Mudanças desde a última exportação ({len(mudancas)}):",
        *[f"  - {m}" for m in mudancas],
    ]
    (saida / "relatorio.txt").write_text("\n".join(relatorio) + "\n", encoding="utf-8")
    print("\n".join(relatorio[:5]))
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    sys.exit(main(*sys.argv[1:]))
